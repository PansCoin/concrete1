# -*- coding: utf-8 -*-
"""note.py - Word (.docx) writer for the explanatory note."""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .engine import (S, Ex, Num, equation_xml, fmt, math_paragraph_xml,
                     name_omml, om_rel, om_text, _r)


class Note:
    def __init__(self, title="Explanatory note"):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(3)
        for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.size = Pt(size)
            hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        self.fig_no = 0
        self.tab_no = 0
        self.tmpdir = None
        self._add_page_numbers()

    # ------------------------------------------------------------------ text
    def h1(self, t, new_page=True):
        if new_page and len(self.doc.paragraphs) > 3:
            self.page_break()
        return self.doc.add_heading(t, level=1)

    def h2(self, t):
        return self.doc.add_heading(t, level=2)

    def h3(self, t):
        return self.doc.add_heading(t, level=3)

    def p(self, text="", bold=False, italic=False, size=None, align=None, color=None):
        par = self.doc.add_paragraph()
        if text:
            r = par.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def bullet(self, text):
        return self.doc.add_paragraph(text, style="List Bullet")

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def ok(self, text, good=True):
        par = self.doc.add_paragraph()
        r = par.add_run(("✔ " if good else "✘ ") + text)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x7F, 0x3B) if good else RGBColor(0xB0, 0x1E, 0x1E)
        return par

    def info(self, text):
        par = self.doc.add_paragraph()
        r = par.add_run("➜ " + text)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x5F, 0xA8)
        return par

    # ------------------------------------------------------------- equations
    def _math(self, inner, desc=None):
        if desc:
            par = self.doc.add_paragraph()
            r = par.add_run(desc)
            r.italic = True
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.keep_with_next = True
        par = self.doc.add_paragraph()
        par._p.append(parse_xml(math_paragraph_xml(inner)))
        par.paragraph_format.space_after = Pt(4)
        return par

    def calc(self, name, expr, unit="", desc=None, show_sym=True, digits=None,
             value=None) -> S:
        """Write 'name = formula = numbers = result unit' and return S(name)."""
        if not isinstance(expr, Ex):
            expr = Num(expr)
        val = expr.value() if value is None else value
        xml = equation_xml(name, expr, val, unit, show_sym=show_sym, digits=digits)
        plain = re.sub(r"<[^>]+>", "", xml)
        if len(plain) > 88 and show_sym and not isinstance(expr, (Num, S)):
            self._math(name_omml(name) + _r("=", True) + expr.om(False), desc)
            self._math(_r("=", True) + expr.om(True) + _r("=", True) + _r(fmt(val, digits), True)
                       + (_r(" " + unit, True) if unit else ""))
        else:
            self._math(xml, desc)
        return S(name, val, unit, digits)

    def given(self, name, value, unit="", desc=None, digits=None) -> S:
        self._math(equation_xml(name, None, value, unit, digits=digits), desc)
        return S(name, value, unit, digits)

    def eq(self, name, expr, unit="", desc=None):
        """Formula only (symbolic), no numbers."""
        inner = name_omml(name) + _r("=", True) + expr.om(False)
        return self._math(inner, desc)

    def compare(self, a: S, op: str, b: S, desc=None, ok_text=None, ng_text=None,
                good=None, neutral=False):
        """a  op  b  (e.g. V_Ed = 45 kN < V_Rd,c = 60 kN)."""
        av, bv = float(a.value()), float(b.value())

        def side(x, v):
            if x.name.replace(".", "").replace(",", "").isdigit():
                return _r(fmt(v, x.digits), True)
            return (name_omml(x.name) + _r("=", True) + _r(fmt(v, x.digits), True) +
                    (_r(" " + x.unit, True) if x.unit else ""))
        if good is None:
            good = {"<": av < bv, "≤": av <= bv + 1e-9, ">": av > bv,
                    "≥": av >= bv - 1e-9}[op]
        inner = (side(a, av) + om_rel(op if good else
                 {"<": "≥", "≤": ">", ">": "≤", "≥": "<"}[op]) + side(b, bv))
        self._math(inner, desc)
        if ok_text or ng_text:
            if neutral and not good:
                self.info(ng_text or "")
            else:
                self.ok(ok_text if good else (ng_text or "NOT satisfied"), good)
        return good

    def math_raw(self, inner, desc=None):
        return self._math(inner, desc)

    # ---------------------------------------------------------- tables/figs
    def table(self, header, rows, caption=None, widths=None, font=9):
        if caption:
            self.tab_no += 1
            cp = self.p(f"Table {self.tab_no}. {caption}", bold=True, size=9.5)
            cp.paragraph_format.keep_with_next = True
        noheader = all(not str(h) for h in header)
        t = self.doc.add_table(rows=0 if noheader else 1, cols=len(header))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate([] if noheader else header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(h))
            r.bold = True
            r.font.size = Pt(font)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                txt = fmt(v) if isinstance(v, float) else str(v)
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(txt)
                r.font.size = Pt(font)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def fig(self, path, caption=None, width_cm=16.0):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(path, width=Cm(width_cm))
        par.paragraph_format.keep_with_next = bool(caption)
        if caption:
            self.fig_no += 1
            cp = self.p(f"Figure {self.fig_no}. {caption}", italic=True, size=9,
                        align="center")
        return par

    # ---------------------------------------------------------------- misc
    def toc(self):
        par = self.doc.add_paragraph()
        run = par.add_run()
        fld = parse_xml(
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:instr="TOC \\o &quot;1-2&quot; \\h \\z \\u"><w:r><w:t>'
            'Right-click here and choose "Update field" to build the table of contents.'
            '</w:t></w:r></w:fldSimple>')
        par._p.append(fld)

    def _add_page_numbers(self):
        footer = self.doc.sections[0].footer
        par = footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par._p.append(parse_xml(
            '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>'))

    def save(self, path):
        self.doc.save(path)
